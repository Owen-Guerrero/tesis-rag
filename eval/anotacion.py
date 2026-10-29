"""
Anotación manual a ciegas de las respuestas + acuerdo entre anotadores (kappa de Cohen)
--------------------------------------------------------------------------------------
1) Preparar (una vez):
     python eval\\anotacion.py --preparar
   Crea eval/anotacion_owen.xlsx y eval/anotacion_irwin.xlsx con las 40 respuestas
   (20 preguntas x 2 arquitecturas) MEZCLADAS y SIN decir de qué arquitectura es cada una.
   La correspondencia queda en eval/anotacion_clave.json (no abrirlo antes de anotar).

2) Cada uno califica SU archivo, sin ver el del otro:
     - correccion:  correcta | parcial | incorrecta     (frente a la respuesta de referencia)
     - fidelidad:   si | no   ("si" = todo lo que afirma está respaldado por la norma;
                               "no" = agrega algo que no está en la referencia ni en la norma)

3) Calcular (cuando los dos terminaron):
     python eval\\anotacion.py --kappa
   Da el kappa de Cohen por criterio y los resultados por arquitectura
   (logs/evaluacion/anotacion.md). Los desacuerdos se resuelven después conversando,
   pero el kappa se reporta con las calificaciones independientes.
"""

import argparse
import json
import random
import sys
from collections import Counter
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipelines.comun import config, leer_jsonl, ruta, slug  # noqa: E402

ANOTADORES = ("owen", "irwin")
CLAVE = ruta("eval/anotacion_clave.json")
OPCIONES = {"correccion": ["correcta", "parcial", "incorrecta"], "fidelidad": ["si", "no"]}


def ultima(arq: str, maquina: str, ids: set) -> dict:
    """Primera corrida del PILOTO (se ignoran archivos de otras preguntas, p. ej. las S01-S05 del OE2)."""
    llm = slug(config()["llm"]["modelo"])
    for archivo in sorted(ruta("resultados").glob(f"{arq}__*__{llm}__{maquina}__*.jsonl")):
        filas = {x["id"]: x for x in leer_jsonl(archivo) if "pregunta" in x}
        if ids <= set(filas):
            return filas
    raise SystemExit(f"No hay resultados del piloto para {arq} en resultados/.")


def preparar(maquina: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.datavalidation import DataValidation

    preguntas = {q["id"]: q for q in leer_jsonl(ruta("eval/preguntas_piloto.jsonl"))}
    items = []
    for arq in ("naive", "raptor"):
        res = ultima(arq, maquina, set(preguntas))
        for i in sorted(preguntas):
            items.append({"arquitectura": arq, "pregunta_id": i, "respuesta": res[i].get("respuesta", "")})
    random.Random(2026).shuffle(items)
    clave = {}
    for n, it in enumerate(items, start=1):
        it["codigo"] = f"R{n:02d}"
        clave[it["codigo"]] = {"arquitectura": it["arquitectura"], "pregunta_id": it["pregunta_id"]}
    CLAVE.write_text(json.dumps(clave, ensure_ascii=False, indent=1), encoding="utf-8")

    for anotador in ANOTADORES:
        wb = Workbook()
        ws = wb.active
        ws.title = "Anotación"
        ws.append(["codigo", "pregunta", "respuesta de referencia", "respuesta a calificar",
                   "correccion", "fidelidad", "comentario"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for it in items:
            q = preguntas[it["pregunta_id"]]
            ws.append([it["codigo"], q["pregunta"], q["respuesta_referencia"], it["respuesta"], "", "", ""])
        for col, ancho in zip("ABCDEFG", [8, 40, 60, 70, 13, 11, 30]):
            ws.column_dimensions[col].width = ancho
        for fila in ws.iter_rows(min_row=2):
            for c in fila:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        for col, campo in (("E", "correccion"), ("F", "fidelidad")):
            dv = DataValidation(type="list", formula1='"' + ",".join(OPCIONES[campo]) + '"', allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}{len(items) + 1}")
        ws.freeze_panes = "B2"
        g = wb.create_sheet("Criterios")
        for t in ["correccion: correcta = responde la pregunta y coincide con lo esencial de la referencia;",
                  "            parcial = acierta en parte, pero omite algo importante o mezcla algo irrelevante;",
                  "            incorrecta = no responde, se equivoca en lo esencial o dice 'No se encontró evidencia'.",
                  "fidelidad:  si = todo lo que afirma está en la referencia o en la norma;",
                  "            no = afirma al menos una cosa que no está (o atribuye un requisito a otro control).",
                  "Califica cada respuesta sola, sin buscar de qué sistema es, y sin ver el archivo del otro anotador."]:
            g.append([t])
        g.column_dimensions["A"].width = 120
        salida = ruta(f"eval/anotacion_{anotador}.xlsx")
        wb.save(salida)
        print(f"Creado {salida.relative_to(ruta('.'))} ({len(items)} respuestas)")
    print(f"Clave guardada en {CLAVE.relative_to(ruta('.'))}: no la abran antes de terminar de anotar.")


def leer(anotador: str) -> dict:
    from openpyxl import load_workbook

    ws = load_workbook(ruta(f"eval/anotacion_{anotador}.xlsx"), read_only=True)["Anotación"]
    datos = {}
    for fila in list(ws.iter_rows(values_only=True))[1:]:
        if fila and fila[0]:
            datos[fila[0]] = {"correccion": (fila[4] or "").strip().lower(),
                              "fidelidad": (fila[5] or "").strip().lower()}
    return datos


def kappa(a: list, b: list) -> float:
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(a) | set(b)) / (n * n)
    return (po - pe) / (1 - pe) if pe < 1 else 1.0


def calcular() -> None:
    clave = json.loads(CLAVE.read_text(encoding="utf-8"))
    A, B = (leer(x) for x in ANOTADORES)
    md = ["# Anotación manual a ciegas (2 anotadores, 40 respuestas)", ""]
    for campo in OPCIONES:
        cod = [c for c in clave if A.get(c, {}).get(campo) and B.get(c, {}).get(campo)]
        if not cod:
            md.append(f"- {campo}: sin datos completos.")
            continue
        a = [A[c][campo] for c in cod]
        b = [B[c][campo] for c in cod]
        acuerdo = sum(x == y for x, y in zip(a, b)) / len(cod)
        md.append(f"- **{campo}**: kappa de Cohen = {kappa(a, b):.3f} | acuerdo observado = {acuerdo:.1%} "
                  f"| respuestas con ambas calificaciones: {len(cod)}/40")
    md += ["", "| Arquitectura | Correctas | Parciales | Incorrectas | Fieles (sí) | Anotador |", "|---|---|---|---|---|---|"]
    for arq in ("naive", "raptor"):
        for nombre, datos in zip(ANOTADORES, (A, B)):
            cs = [datos[c] for c in clave if clave[c]["arquitectura"] == arq and c in datos]
            cc = Counter(x["correccion"] for x in cs)
            fi = sum(x["fidelidad"] == "si" for x in cs)
            md.append(f"| {arq} | {cc['correcta']} | {cc['parcial']} | {cc['incorrecta']} | {fi}/{len(cs)} | {nombre} |")
    salida = ruta("logs/evaluacion/anotacion.md")
    salida.parent.mkdir(parents=True, exist_ok=True)
    salida.write_text("\n".join(md), encoding="utf-8")
    print("\n".join(md))
    print(f"\nGuardado: {salida.relative_to(ruta('.'))}")


if __name__ == "__main__":
    p = argparse.ArgumentParser()
    p.add_argument("--preparar", action="store_true")
    p.add_argument("--kappa", action="store_true")
    p.add_argument("--maquina", default="lab_p5")
    a = p.parse_args()
    if a.preparar:
        preparar(a.maquina)
    elif a.kappa:
        calcular()
    else:
        print(__doc__)

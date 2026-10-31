"""
Contraste entre el juez automático (RAGAS) y la anotación humana a ciegas
-------------------------------------------------------------------------
Responde: ¿cuánto se parece la calificación del LLM juez a la de los anotadores humanos
sobre LAS MISMAS 40 respuestas (20 preguntas x 2 arquitectura, 1.ª corrida del piloto)?

Insumos (todos en el repositorio, no requiere GPU ni el laboratorio):
  eval/anotacion_clave.json                 código R01..R40 -> arquitectura y pregunta
  eval/anotacion_<anotador>.xlsx            una planilla por anotador (ver eval/anotacion.py)
  logs/evaluacion/ragas_por_pregunta.csv    salida de eval/ragas_eval.py

Escala humana de corrección: correcta = 1, parcial = 0,5, incorrecta = 0
(media de los dos anotadores). Se compara con Answer Correctness y Answer Similarity
de RAGAS mediante la correlación de Spearman (rho), con IC 95 % bootstrap (2000 iteraciones).

Uso:
  python eval\\contraste_ragas_humano.py
Salida: logs/evaluacion/contraste_ragas_humano.md
"""

import csv
import json
import random
import statistics
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipelines.comun import ruta  # noqa: E402

ANOTADORES = ("owen", "irwin")   # nombres de las planillas eval/anotacion_<nombre>.xlsx
ESCALA = {"correcta": 1.0, "parcial": 0.5, "incorrecta": 0.0}
FIDELIDAD = {"si": 1.0, "no": 0.0}


def leer_anotacion(nombre: str) -> dict:
    from openpyxl import load_workbook

    ws = load_workbook(ruta(f"eval/anotacion_{nombre}.xlsx"), read_only=True)["Anotación"]
    datos = {}
    for fila in list(ws.iter_rows(values_only=True))[1:]:
        if fila and fila[0]:
            datos[fila[0]] = {"correccion": str(fila[4] or "").strip().lower(),
                              "fidelidad": str(fila[5] or "").strip().lower()}
    return datos


def rangos(v: list) -> list:
    """Rangos promedio (maneja empates)."""
    orden = sorted(range(len(v)), key=lambda i: v[i])
    r = [0.0] * len(v)
    i = 0
    while i < len(v):
        j = i
        while j + 1 < len(v) and v[orden[j + 1]] == v[orden[i]]:
            j += 1
        for k in range(i, j + 1):
            r[orden[k]] = (i + j) / 2 + 1
        i = j + 1
    return r


def pearson(x: list, y: list):
    if len(x) < 3 or statistics.pstdev(x) == 0 or statistics.pstdev(y) == 0:
        return None
    mx, my = statistics.mean(x), statistics.mean(y)
    num = sum((a - mx) * (b - my) for a, b in zip(x, y))
    den = (sum((a - mx) ** 2 for a in x) * sum((b - my) ** 2 for b in y)) ** 0.5
    return num / den


def spearman(x: list, y: list):
    return pearson(rangos(x), rangos(y))


def ic_bootstrap(x: list, y: list, n=2000, semilla=42):
    rng = random.Random(semilla)
    idx = list(range(len(x)))
    valores = []
    for _ in range(n):
        m = rng.choices(idx, k=len(idx))
        r = spearman([x[i] for i in m], [y[i] for i in m])
        if r is not None:
            valores.append(r)
    if len(valores) < n * 0.5:
        return None, None
    valores.sort()
    return valores[int(0.025 * len(valores))], valores[int(0.975 * len(valores)) - 1]


def main() -> None:
    clave = json.loads(ruta("eval/anotacion_clave.json").read_text(encoding="utf-8"))
    a1, a2 = (leer_anotacion(x) for x in ANOTADORES)
    with open(ruta("logs/evaluacion/ragas_por_pregunta.csv"), encoding="utf-8") as fh:
        ragas = {(r["arquitectura"], r["id"]): r for r in csv.DictReader(fh)}

    filas = []
    for cod, c in sorted(clave.items()):
        r = ragas.get((c["arquitectura"], c["pregunta_id"]))
        a, b = a1.get(cod, {}), a2.get(cod, {})
        if not r or a.get("correccion") not in ESCALA or b.get("correccion") not in ESCALA:
            continue
        num = lambda k: float(r[k]) if r.get(k) not in (None, "", "None") else None  # noqa: E731
        filas.append({"codigo": cod, "arquitectura": c["arquitectura"], "id": c["pregunta_id"],
                      "humano": (ESCALA[a["correccion"]] + ESCALA[b["correccion"]]) / 2,
                      "anotador1": ESCALA[a["correccion"]], "anotador2": ESCALA[b["correccion"]],
                      "consenso": a["correccion"] if a["correccion"] == b["correccion"] else None,
                      "fid_humana": (FIDELIDAD.get(a.get("fidelidad"), 0) + FIDELIDAD.get(b.get("fidelidad"), 0)) / 2,
                      "answer_correctness": num("answer_correctness"),
                      "semantic_similarity": num("semantic_similarity"),
                      "faithfulness": num("faithfulness")})
    if not filas:
        raise SystemExit("No se pudo emparejar ninguna respuesta (revisa la clave y ragas_por_pregunta.csv).")

    md = ["# Contraste entre el juez automático (RAGAS) y la anotación humana", "",
          f"Respuestas emparejadas: {len(filas)} (misma 1.ª corrida del piloto). "
          "Escala humana: correcta = 1, parcial = 0,5, incorrecta = 0. "
          "Spearman rho [IC 95 % bootstrap, 2000 iteraciones].", "",
          "| Métrica RAGAS | vs. media de anotadores | vs. anotador 1 | vs. anotador 2 | n |", "|---|---|---|---|---|"]
    for m, nombre in (("answer_correctness", "Answer Correctness"), ("semantic_similarity", "Answer Similarity")):
        par = [f for f in filas if f[m] is not None]
        celdas = []
        for h in ("humano", "anotador1", "anotador2"):
            x, y = [f[m] for f in par], [f[h] for f in par]
            rho = spearman(x, y)
            if rho is None:
                celdas.append("—")
                continue
            lo, hi = ic_bootstrap(x, y)
            celdas.append(f"{rho:.3f}" + (f" [{lo:.3f}, {hi:.3f}]" if lo is not None else ""))
        md.append(f"| {nombre} | {celdas[0]} | {celdas[1]} | {celdas[2]} | {len(par)} |")

    par = [f for f in filas if f["faithfulness"] is not None]
    rho_f = spearman([f["faithfulness"] for f in par], [f["fid_humana"] for f in par])
    md += ["", f"Faithfulness (RAGAS) vs. fidelidad humana: rho = {rho_f:.3f} (n = {len(par)})."
           if rho_f is not None else
           f"Faithfulness (RAGAS) vs. fidelidad humana: no calculable (la fidelidad humana casi no varía; n = {len(par)}).",
           "Nota: el criterio humano de fidelidad resultó demasiado estricto (ver anotacion.md), "
           "por lo que este contraste se interpreta con cautela.", "",
           "## Answer Correctness (RAGAS) según la calificación humana de consenso", "",
           "| Calificación (ambos anotadores coinciden) | n | Answer Correctness media | Faithfulness media |",
           "|---|---|---|---|"]
    for cat in ("correcta", "parcial", "incorrecta"):
        g = [f for f in filas if f["consenso"] == cat]
        ac = [f["answer_correctness"] for f in g if f["answer_correctness"] is not None]
        fa = [f["faithfulness"] for f in g if f["faithfulness"] is not None]
        md.append(f"| {cat} | {len(g)} | {statistics.mean(ac):.3f} | {statistics.mean(fa):.3f} |"
                  if ac and fa else f"| {cat} | {len(g)} | — | — |")
    sin = sum(f["consenso"] is None for f in filas)
    md += [f"| sin consenso | {sin} | | |", "",
           "Lectura: rho cercano a 1 = el juez ordena las respuestas como los humanos; "
           "cercano a 0 = no hay relación. Con n = 40 los intervalos son amplios."]

    salida = ruta("logs/evaluacion/contraste_ragas_humano.md")
    salida.write_text("\n".join(md), encoding="utf-8")
    print("\n".join(md))
    print(f"\nGuardado: {salida.relative_to(ruta('.'))}")


if __name__ == "__main__":
    main()

"""
Revisar el árbol RAPTOR (resúmenes) y preparar la auditoría de alucinaciones
--------------------------------------------------------------------------
Uso (desde la raíz del repo):
  python pipelines\\ver_arbol.py                       resumen del árbol por capa
  python pipelines\\ver_arbol.py --capa 3              muestra los resúmenes de la capa 3
  python pipelines\\ver_arbol.py --id raptor-c2-5      un resumen + el texto de sus hijos
  python pipelines\\ver_arbol.py --auditoria 30        muestra estratificada de 30 resúmenes
                                                      -> corpus/procesado/arbol_auditoria.xlsx

La auditoría compara cada resumen con el texto de sus hijos para detectar
afirmaciones que no están en ellos. El Excel queda en corpus/procesado/, que
NO se sube a GitHub (contiene texto ISO). Lo que se reporta es el conteo.
"""

import argparse
import json
import random
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pipelines.comun import config, modelo_embeddings, ruta, slug  # noqa: E402


def cargar_arbol(prueba: int = None) -> dict:
    sufijo = f"{slug(modelo_embeddings())}__{slug(config()['llm']['modelo'])}" + (f"__prueba{prueba}" if prueba else "")
    archivo = ruta("corpus/procesado") / f"arbol_raptor__{sufijo}.jsonl"
    if not archivo.exists():
        raise SystemExit(f"No existe {archivo.relative_to(ruta('.'))}: corre antes pipelines\\raptor_arbol.py")
    nodos = {}
    for linea in archivo.read_text(encoding="utf-8").splitlines():
        n = json.loads(linea)
        nodos[n["id"]] = n
    return nodos


def mostrar(n: dict, nodos: dict, con_hijos: bool) -> None:
    print(f"\n===== {n['id']} | capa {n['nivel']} | {n['n_tokens']} tokens | {len(n['hijos'])} hijos =====")
    print(f"Citas: {' '.join(n.get('citas', []))}")
    print(n["texto"])
    if con_hijos:
        for h in n["hijos"]:
            hijo = nodos[h]
            print(f"\n  --- hijo {h} (capa {hijo['nivel']}) ---\n  " + hijo["texto"].replace("\n", "\n  "))


def auditoria(nodos: dict, total: int) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    por_capa = defaultdict(list)
    for n in nodos.values():
        if n["nivel"] > 0:
            por_capa[n["nivel"]].append(n)
    resumenes = sum(len(v) for v in por_capa.values())
    rng = random.Random(42)
    muestra = []
    for capa, lista in sorted(por_capa.items()):
        k = max(2, round(total * len(lista) / resumenes))      # al menos 2 por capa
        muestra += rng.sample(lista, min(k, len(lista)))
    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoría"
    cab = ["id", "capa", "tokens", "resumen", "texto de los hijos", "afirmaciones revisadas",
           "afirmaciones NO respaldadas por los hijos", "¿identificadores correctos? (sí/no)", "comentario", "revisor"]
    ws.append(cab)
    for c in ws[1]:
        c.font = Font(bold=True)
    for n in muestra:
        hijos = "\n\n".join(f"[{h}] {nodos[h]['texto']}" for h in n["hijos"])
        ws.append([n["id"], n["nivel"], n["n_tokens"], n["texto"], hijos[:32000], "", "", "", "", ""])
    for col, ancho in zip("ABCDEFGHIJ", [16, 6, 8, 70, 90, 12, 18, 14, 30, 10]):
        ws.column_dimensions[col].width = ancho
    for fila in ws.iter_rows(min_row=2):
        for c in fila:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    salida = ruta("corpus/procesado") / "arbol_auditoria.xlsx"
    wb.save(salida)
    print(f"Muestra de {len(muestra)} resúmenes ({', '.join(f'capa {c}: {sum(m['nivel'] == c for m in muestra)}' for c in sorted(por_capa))})")
    print(f"Guardado en {salida.relative_to(ruta('.'))} (no se sube a GitHub)")


def main() -> None:
    p = argparse.ArgumentParser()
    p.add_argument("--capa", type=int)
    p.add_argument("--id")
    p.add_argument("--auditoria", type=int, metavar="N")
    p.add_argument("--prueba", type=int, help="árbol de prueba (p. ej. 300)")
    p.add_argument("--max", type=int, default=3, help="cuántos resúmenes mostrar con --capa")
    a = p.parse_args()
    nodos = cargar_arbol(a.prueba)

    if a.id:
        mostrar(nodos[a.id], nodos, con_hijos=True)
    elif a.capa:
        for n in [n for n in nodos.values() if n["nivel"] == a.capa][:a.max]:
            mostrar(n, nodos, con_hijos=False)
    elif a.auditoria:
        auditoria(nodos, a.auditoria)
    else:
        capas = defaultdict(list)
        for n in nodos.values():
            capas[n["nivel"]].append(n["n_tokens"])
        for c, v in sorted(capas.items()):
            nombre = "hojas" if c == 0 else "resúmenes"
            print(f"capa {c}: {len(v)} {nombre} | media {sum(v) / len(v):.0f} tokens | máx {max(v)}")


if __name__ == "__main__":
    main()
